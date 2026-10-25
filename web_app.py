
from pathlib import Path
import io
import re
from functools import lru_cache


import joblib
import pandas as pd
from docx import Document
from flask import Flask, jsonify, render_template, request
from time import perf_counter
from src.skill_matcher import detect_skills, analyze_skill_gap
from src.recommender import recommend_skills
from pypdf import PdfReader

BASE_DIR = Path(__file__).resolve().parent
DATASET_DIR = BASE_DIR / "dataset"
MODEL_DIR = BASE_DIR / "models"

SKILL_FILE = DATASET_DIR / "skill_catalog.csv"
MODEL_FILE = MODEL_DIR / "rejection_model.joblib"
FEATURES_FILE = MODEL_DIR / "features.joblib"

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # 10 MB


@lru_cache(maxsize=1)
def load_catalog():
    if not SKILL_FILE.exists():
        return pd.DataFrame(
            columns=[
                "skill",
                "category",
                "priority",
                "learning_suggestion",
            ]
        )

    df = pd.read_csv(SKILL_FILE).fillna("")

    for col in [
        "skill",
        "category",
        "priority",
        "learning_suggestion",
    ]:
        if col not in df.columns:
            df[col] = ""

    df["skill"] = df["skill"].astype(str).str.strip()
    return df[df["skill"] != ""]


def extract_resume_text(upload):
    name = upload.filename.lower()
    content = upload.read()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)

        return "\n".join(parts)

    raise ValueError("Upload a PDF or DOCX resume.")


def find_skills(text, skill_names):
    found = []
    text = text.lower()

    for skill in skill_names:
        skill = str(skill).strip()
        if not skill:
            continue

        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(skill.lower())
            + r"(?![a-z0-9])"
        )

        if re.search(pattern, text):
            found.append(skill)

    return found


@lru_cache(maxsize=1)
def get_saved_model():
    if not MODEL_FILE.exists() or not FEATURES_FILE.exists():
        return None, None

    try:
        return (
            joblib.load(MODEL_FILE),
            joblib.load(FEATURES_FILE),
        )
    except Exception:
        return None, None



# Requirement analysis is evidence-based keyword/phrase matching, not a guarantee
# that a candidate has or lacks a skill. Requirements are inferred from the JD.
REQUIREMENT_GROUPS = {
    "Education & qualifications": [
        "bachelor", "bachelor's", "b.tech", "btech", "degree", "master", "master's",
        "phd", "certification", "certified", "qualification", "computer science",
        "artificial intelligence", "engineering"
    ],
    "Experience & seniority": [
        "years of experience", "year experience", "experience in", "work experience",
        "professional experience", "senior", "lead", "manager", "management",
        "entry level", "fresher", "internship"
    ],
    "Responsibilities & delivery": [
        "responsible for", "responsibilities", "deliver", "deployment", "production",
        "stakeholder", "business impact", "requirements gathering", "problem solving",
        "cross-functional", "end-to-end", "project delivery"
    ],
    "Leadership & communication": [
        "leadership", "team management", "manage a team", "mentoring", "mentor",
        "communication", "collaboration", "presentation", "lead a team"
    ],
    "Projects & practical evidence": [
        "projects", "portfolio", "hands-on", "practical experience", "case study",
        "implementation", "built", "developed", "designed", "deployed"
    ],
    "Tools, platforms & technologies": [
        "python", "sql", "mysql", "postgresql", "excel", "power bi", "tableau",
        "pandas", "scikit-learn", "machine learning", "deep learning", "nlp",
        "artificial intelligence", "generative ai", "genai", "llm", "rag", "azure", "aws", "gcp", "docker",
        "kubernetes", "git", "github", "ci/cd", "mlops", "flask", "django",
        "tensorflow", "pytorch", "data visualization", "statistics", "etl", "dax",
        "data modeling", "star schema", "html", "css", "jinja2", "power bi dashboard"
    ],
}

TERM_VARIANTS = {
    "management": ["management", "manage", "manages", "managed", "manager", "managing", "team management"],
    "manager": ["manager", "management", "manage", "managed", "managing"],
    "lead": ["lead", "leads", "led", "leading", "leader", "leadership"],
    "leadership": ["leadership", "leader", "leading", "lead", "led"],
    "deliver": ["deliver", "delivers", "delivered", "delivering", "delivery"],
    "delivery": ["delivery", "deliver", "delivered", "delivering"],
    "deployment": ["deployment", "deploy", "deployed", "deploying"],
    "production": ["production", "production-ready", "in production"],
    "stakeholder": ["stakeholder", "stakeholders", "stakeholder engagement", "stakeholder management"],
    "responsible for": ["responsible for", "responsibilities include", "responsible for delivering"],
    "responsibilities": ["responsibilities", "responsibility", "responsible for"],
    "implementation": ["implementation", "implement", "implemented", "implementing"],
    "built": ["built", "build", "building"],
    "developed": ["developed", "develop", "developing", "development"],
    "designed": ["designed", "design", "designing", "designs"],
    "projects": ["projects", "project"],
    "communication": ["communication", "communicate", "communicated", "communicating"],
    "collaboration": ["collaboration", "collaborate", "collaborated", "collaborating", "collaborative"],
    "internship": ["internship", "intern", "interned", "internship program"],
    "experience in": ["experience in", "experienced in", "experience with"],
}

def _contains_phrase(text, phrase):
    text = re.sub(r"[^a-z0-9+#./ -]", " ", text.lower())
    phrase = phrase.lower()
    candidates = TERM_VARIANTS.get(phrase, [phrase])
    for candidate in candidates:
        if len(candidate) <= 3 and candidate not in {"aws", "gcp", "nlp", "llm", "rag", "etl"}:
            found = re.search(r"(?<![a-z0-9])" + re.escape(candidate) + r"(?![a-z0-9])", text)
        else:
            found = re.search(r"(?<![a-z0-9])" + re.escape(candidate) + r"(?![a-z0-9])", text)
        if found:
            return True
    return False

def analyze_requirements(resume_text, job_description):
    resume = re.sub(r"\s+", " ", resume_text.casefold())
    jd = re.sub(r"\s+", " ", job_description.casefold())
    categories = []
    reasons = []
    total = matched_count = partial_count = missing_count = 0

    for category, candidates in REQUIREMENT_GROUPS.items():
        required = [term for term in candidates if _contains_phrase(jd, term)]
        if not required:
            continue
        # Keep two separate views: (1) JD requirements evidenced by the resume,
        # which drive the status and coverage score; (2) all tracked terms found
        # in the resume for this category, which are useful context for the user.
        detected = [term for term in required if _contains_phrase(resume, term)]
        detected_in_resume = [term for term in candidates if _contains_phrase(resume, term)]
        # A partial status means some JD terms in this category are evidenced,
        # but other detected JD terms are not; this is a transparent heuristic.
        if len(detected) == len(required):
            status = "Matched"
        elif detected:
            status = "Partially matched"
        else:
            status = "Not detected"
        total += len(required)
        matched_count += len(detected)
        partial_count += int(status == "Partially matched")
        missing_count += int(status == "Not detected")
        categories.append({
            "category": category,
            "status": status,
            "detected": detected,
            "detected_in_resume": detected_in_resume,
            "requirements_review": [term for term in required if term not in detected],
            "required_count": len(required),
            "detected_count": len(detected),
            "explanation": (
                f"Detected {len(detected)} of {len(required)} tracked terms from the job description."
            )
        })
        if status != "Matched":
            missing_terms = [term for term in required if term not in detected]
            reasons.append({
                "category": category,
                "reason": (
                    "The resume does not clearly demonstrate all tracked requirements in this area."
                    if detected else
                    "No clear evidence for the tracked requirements in this area was detected in the resume text."
                ),
                "requirements_to_review": missing_terms,
                "priority": "High" if category in {
                    "Experience & seniority", "Tools, platforms & technologies",
                    "Education & qualifications", "Responsibilities & delivery"
                } else "Medium",
                "improvement": (
                    "If you have this experience, add a concise example with your role, tools used, and outcome. "
                    "If not, consider a relevant project, course, or practical exercise."
                )
            })

    return {
        "categories": categories,
        "potential_rejection_reasons": reasons,
        "tracked_requirement_terms": total,
        "detected_requirement_terms": matched_count,
        "requirement_coverage_percent": round(100 * matched_count / total, 1) if total else 0,
        "note": (
            "These are possible gaps inferred from the job description, not the employer's confirmed reason. "
            "A term not found in the resume may mean it is omitted or phrased differently, not that the candidate lacks it."
        )
    }


@app.get("/")
def home():
    return render_template("index.html")


@app.get("/api/health")
def health():
    return jsonify({"status": "ok"})


@app.post("/api/analyze")
def analyze():
    request_started = perf_counter()
    upload = request.files.get("resume")
    job_description = request.form.get(
        "job_description", ""
    ).strip()

    try:
        experience = float(
            request.form.get("experience", "0")
        )
    except ValueError:
        return jsonify({"error": "Experience must be a number."}), 400

    if not upload or not upload.filename:
        return jsonify({"error": "Please upload your resume."}), 400

    if not upload.filename.lower().endswith((".pdf", ".docx")):
        return jsonify({"error": "Only PDF and DOCX files are supported."}), 400

    if not job_description:
        return jsonify({"error": "Please enter the job description."}), 400

    parse_started = perf_counter()
    try:
        resume_text = extract_resume_text(upload)
    except Exception as exc:
        return jsonify({"error": f"Could not read resume: {exc}"}), 400

    parse_seconds = perf_counter() - parse_started

    if not resume_text.strip():
        return jsonify({
            "error": "No readable text found. Scanned PDFs may require OCR."
        }), 400

    catalog = load_catalog()
    if catalog.empty:
        return jsonify({
            "error": "skill_catalog.csv is missing or empty in dataset/."
        }), 500

    # Skill detection is independent of the recommendation catalog. Include catalog labels
    # as additional supported terms, while canonical aliases handle GenAI/Generative AI.
    extra_skills = catalog["skill"].tolist() if not catalog.empty else []
    resume_skills = detect_skills(resume_text, extra_skills)
    required_skills = detect_skills(job_description, extra_skills)
    skill_gap = analyze_skill_gap(resume_skills, required_skills)
    matched = skill_gap["matched_skills"]
    missing = skill_gap["missing_skills"]
    ratio = (len(matched) / len(required_skills)) if required_skills else 0.0
    requirement_analysis = analyze_requirements(resume_text, job_description)

    # Do not present a tiny catalog sample as an overall resume-to-job score.
    # Require at least three tracked JD terms before showing a percentage.
    overall_match_percent = (
        requirement_analysis["requirement_coverage_percent"]
        if requirement_analysis["tracked_requirement_terms"] >= 3
        else None
    )
    # These are tracked job-description terms not detected in the resume text.
    # They are review prompts, not proof the candidate lacks the skills.
    potential_gaps = list(dict.fromkeys(
        term
        for category in requirement_analysis["categories"]
        for term in category["requirements_review"]
    ))
    potential_gap_count = len(potential_gaps)

    # Surface matched requirement terms alongside canonical skills. Category cards
    # and the matched panel now share the same evidence, while remaining distinct
    # metrics: skills are canonicalized; requirements are job-description phrases.
    matched_requirement_terms = []
    for category in requirement_analysis["categories"]:
        for term in category["detected"]:
            item = {"category": category["category"], "term": term}
            if item not in matched_requirement_terms:
                matched_requirement_terms.append(item)
    matched_items = list(dict.fromkeys(
        matched + [f"{item['category']}: {item['term']}" for item in matched_requirement_terms]
    ))

    # Roadmap covers missing canonical skills plus requirement terms not detected.
    # Generic non-skill phrases are kept in gaps but can receive fallback advice too.
    roadmap_gaps = list(dict.fromkeys(missing + potential_gaps))
    recommendations = recommend_skills(roadmap_gaps, SKILL_FILE)

    # The saved model expects the features used in train_model.py.
    # Its labels come from synthetic demonstration data.
    model_result = {
        "available": False,
        "label": None,
        "scores": [],
        "notice": (
            "Model output is experimental. It was trained on synthetic "
            "demonstration data and is not a validated rejection probability."
        ),
    }

    model, features = get_saved_model()

    if model is not None and features is not None:
        try:
            feature_values = {
                "experience_years": experience,
                "skills_match_ratio": ratio,
            }

            if all(feature in feature_values for feature in features):
                model_input = pd.DataFrame([{
                    feature: feature_values[feature]
                    for feature in features
                }])

                prediction = model.predict(model_input)[0]
                model_result["available"] = True
                model_result["label"] = str(prediction)

                if hasattr(model, "predict_proba"):
                    probabilities = model.predict_proba(model_input)[0]
                    model_result["scores"] = [
                        {
                            "class": str(label),
                            "score": round(float(score) * 100, 1),
                        }
                        for label, score in zip(
                            model.classes_, probabilities
                        )
                    ]
            else:
                model_result["notice"] += (
                    " Saved model features do not match the app inputs."
                )
        except Exception:
            model_result["notice"] += (
                " The saved model could not be evaluated with these inputs."
            )

    return jsonify({
        "filename": upload.filename,
        "experience_years": experience,
        "skill_match_percent": overall_match_percent,
        "catalog_skill_match_percent": round(ratio * 100, 1),
        "potential_gap_count": potential_gap_count,
        "potential_gaps": potential_gaps,
        "required_skill_count": len(required_skills),
        "matched_skill_count": len(matched),
        "matched_requirement_terms": matched_requirement_terms,
        "matched_item_count": len(matched_items),
        "matched_items": matched_items,
        "missing_skill_count": len(missing),
        "analysis_timings": {"resume_parse_seconds": round(parse_seconds, 3), "total_seconds": round(perf_counter() - request_started, 3)},
        "skill_score_warning": (
            "Insufficient job requirements were detected to calculate a meaningful overall percentage. Review the requirement analysis below."
            if overall_match_percent is None else
            "This is a keyword-based requirement coverage estimate, not a hiring probability."
        ),
        "matched_skills": matched,
        "missing_skills": missing,
        "recommendations": recommendations,
        "requirement_analysis": requirement_analysis,
        "model": model_result,
        "summary": (
            "No supported skills were detected in the job description."
            if not required_skills
            else f"{len(matched)} of {len(required_skills)} detected job skills matched."
        ),
    })


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({"error": "File too large. Maximum size is 10 MB."}), 413


if __name__ == "__main__":
    app.run(debug=True, port=5000)
