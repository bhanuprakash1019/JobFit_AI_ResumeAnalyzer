
import re
from pathlib import Path

import joblib
import pandas as pd
import streamlit as st

# Resume parsing libraries
import fitz  # PyMuPDF
from docx import Document


# ---------------- CONFIGURATION ----------------

BASE_DIR = Path(__file__).resolve().parent
DATASET_DIR = BASE_DIR / "dataset"
MODEL_DIR = BASE_DIR / "models"

SKILL_FILE = DATASET_DIR / "skill_catalog.csv"
MODEL_FILE = MODEL_DIR / "rejection_model.joblib"
FEATURES_FILE = MODEL_DIR / "features.joblib"


# ---------------- PAGE SETTINGS ----------------

st.set_page_config(
    page_title="JobFit AI",
    page_icon="📄",
    layout="wide",
)

st.title("📄 JobFit AI")
st.subheader("Resume Analysis & Job Skill Gap Predictor")

st.write(
    "Upload your resume and enter a job description to identify "
    "matching skills, missing technologies, and recommended learning."
)


# ---------------- LOAD SKILL CATALOG ----------------

@st.cache_data
def load_skill_catalog():
    if not SKILL_FILE.exists():
        return pd.DataFrame(
            columns=[
                "skill",
                "category",
                "priority",
                "learning_suggestion",
            ]
        )

    catalog = pd.read_csv(SKILL_FILE)

    # Ensure expected columns exist
    for column in [
        "skill",
        "category",
        "priority",
        "learning_suggestion",
    ]:
        if column not in catalog.columns:
            catalog[column] = ""

    catalog = catalog.fillna("")
    catalog["skill"] = catalog["skill"].astype(str).str.strip()
    return catalog


catalog = load_skill_catalog()


# ---------------- RESUME PARSER ----------------

def extract_resume_text(uploaded_file):
    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):
        text_parts = []

        with fitz.open(
            stream=uploaded_file.getvalue(),
            filetype="pdf",
        ) as pdf:
            for page in pdf:
                text_parts.append(page.get_text())

        return "\n".join(text_parts)

    if file_name.endswith(".docx"):
        import io

        document = Document(
            io.BytesIO(uploaded_file.getvalue())
        )

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        # Include text from tables, which often contain resume details
        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(
                    cell.text for cell in row.cells
                )

        return "\n".join(paragraphs)

    raise ValueError("Please upload a PDF or DOCX file.")


# ---------------- SKILL EXTRACTION ----------------

def find_skills(text, skill_names):
    found = []
    normalized_text = text.lower()

    for skill in skill_names:
        skill = str(skill).strip()

        if not skill:
            continue

        # Match whole terms to reduce false matches
        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(skill.lower())
            + r"(?![a-z0-9])"
        )

        if re.search(pattern, normalized_text):
            found.append(skill)

    return list(dict.fromkeys(found))


def estimate_experience(text):
    """Basic estimate; verify it because resume wording varies."""
    patterns = [
        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s+of\s+experience",
        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s+experience",
    ]

    for pattern in patterns:
        match = re.search(pattern, text.lower())

        if match:
            return float(match.group(1))

    return 0.0


# ---------------- LOAD TRAINED MODEL ----------------

@st.cache_resource
def load_trained_model():
    if not MODEL_FILE.exists():
        return None, None

    try:
        model = joblib.load(MODEL_FILE)

        features = None
        if FEATURES_FILE.exists():
            features = joblib.load(FEATURES_FILE)

        return model, features

    except Exception as error:
        st.warning(f"Could not load the model: {error}")
        return None, None


model, model_features = load_trained_model()


# ---------------- INPUT SECTION ----------------

left, right = st.columns(2)

with left:
    uploaded_resume = st.file_uploader(
        "Upload your resume",
        type=["pdf", "docx"],
        help="Supported formats: PDF and DOCX",
    )

with right:
    experience = st.number_input(
        "Relevant work experience (years)",
        min_value=0.0,
        max_value=50.0,
        value=0.0,
        step=0.5,
        help="Enter your actual relevant experience.",
    )

job_description = st.text_area(
    "Paste the job description",
    height=220,
    placeholder=(
        "Paste the job responsibilities, required skills, "
        "tools, and qualifications here..."
    ),
)

analyze = st.button(
    "Analyze Resume",
    type="primary",
    use_container_width=True,
)


# ---------------- ANALYSIS ----------------

if analyze:
    if uploaded_resume is None:
        st.error("Please upload your resume first.")

    elif not job_description.strip():
        st.error("Please enter the job description.")

    elif catalog.empty:
        st.error(
            "Skill catalog is missing or empty. "
            "Check dataset/skill_catalog.csv."
        )

    else:
        try:
            with st.spinner("Analyzing your resume..."):
                resume_text = extract_resume_text(
                    uploaded_resume
                )

                if not resume_text.strip():
                    st.error(
                        "No readable text was found. "
                        "If this is a scanned PDF, OCR may be required."
                    )
                    st.stop()

                skill_names = catalog["skill"].tolist()

                resume_skills = find_skills(
                    resume_text,
                    skill_names,
                )

                required_skills = find_skills(
                    job_description,
                    skill_names,
                )

                matched_skills = [
                    skill for skill in required_skills
                    if skill.lower() in {
                        item.lower() for item in resume_skills
                    }
                ]

                missing_skills = [
                    skill for skill in required_skills
                    if skill.lower() not in {
                        item.lower() for item in resume_skills
                    }
                ]

                if required_skills:
                    match_ratio = (
                        len(matched_skills) /
                        len(required_skills)
                    )
                else:
                    match_ratio = 0.0

            # ---------------- RESULTS ----------------

            st.divider()
            st.header("Analysis Results")

            metric1, metric2, metric3 = st.columns(3)

            metric1.metric(
                "Skill Match",
                f"{match_ratio * 100:.1f}%",
            )

            metric2.metric(
                "Matched Skills",
                len(matched_skills),
            )

            metric3.metric(
                "Missing Skills",
                len(missing_skills),
            )

            st.progress(match_ratio)

            tab1, tab2, tab3, tab4 = st.tabs(
                [
                    "Matched Skills",
                    "Missing Skills",
                    "Learning Plan",
                    "Model Assessment",
                ]
            )

            with tab1:
                st.subheader("Skills found in both")

                if matched_skills:
                    for skill in matched_skills:
                        st.write(f"✅ {skill}")
                else:
                    st.info(
                        "No matching catalog skills were found."
                    )

                st.caption(
                    "Skill matching is based on text found in "
                    "the resume and job description."
                )

            with tab2:
                st.subheader("Potential skill gaps")

                if missing_skills:
                    for skill in missing_skills:
                        st.write(f"❌ {skill}")
                else:
                    st.success(
                        "No missing catalog skills were detected."
                    )

                st.caption(
                    "A missing skill means its term was not detected. "
                    "It does not prove you lack that skill."
                )

            with tab3:
                st.subheader("Recommended learning")

                if missing_skills:
                    recommendations = catalog[
                        catalog["skill"].str.lower().isin(
                            {
                                skill.lower()
                                for skill in missing_skills
                            }
                        )
                    ]

                    for _, row in recommendations.iterrows():
                        st.markdown(
                            f"**{row['skill']}** "
                            f"— {row['category']}"
                        )

                        if row["priority"]:
                            st.write(
                                f"Priority: {row['priority']}"
                            )

                        if row["learning_suggestion"]:
                            st.write(
                                row["learning_suggestion"]
                            )

                        st.divider()
                else:
                    st.success(
                        "No skill-gap learning suggestions "
                        "are available for this job description."
                    )

            with tab4:
                st.subheader("Experimental model assessment")

                if model is None:
                    st.warning(
                        "Trained model not found. Complete model "
                        "training and save it in the models folder."
                    )

                elif not model_features:
                    st.warning(
                        "The saved feature list is missing. "
                        "Cannot safely construct model input."
                    )

                else:
                    # The current training script uses these features.
                    input_values = {
                        "experience_years": float(experience),
                        "skills_match_ratio": float(match_ratio),
                    }

                    try:
                        model_input = pd.DataFrame(
                            [
                                {
                                    feature: input_values[feature]
                                    for feature in model_features
                                }
                            ]
                        )

                        prediction = model.predict(model_input)[0]

                        st.write(
                            f"Model output label: **{prediction}**"
                        )

                        st.info(
                            "This is an experimental classification "
                            "based on synthetic demonstration data. "
                            "It is not a reliable real-world rejection "
                            "probability or a confirmed rejection reason."
                        )

                        if hasattr(model, "predict_proba"):
                            probabilities = model.predict_proba(
                                model_input
                            )[0]

                            st.write(
                                "Experimental class scores:"
                            )

                            st.dataframe(
                                pd.DataFrame(
                                    {
                                        "Class": model.classes_,
                                        "Score": [
                                            f"{value * 100:.1f}%"
                                            for value in probabilities
                                        ],
                                    }
                                ),
                                use_container_width=True,
                                hide_index=True,
                            )

                            st.caption(
                                "These scores are not validated "
                                "real-world rejection probabilities."
                            )

                    except Exception as error:
                        st.error(
                            "The saved model's expected features "
                            f"do not match the app input: {error}"
                        )

            # ---------------- DOWNLOAD REPORT ----------------

            report = [
                "JOBFIT AI - RESUME ANALYSIS",
                "",
                f"Resume: {uploaded_resume.name}",
                f"Skill match: {match_ratio * 100:.1f}%",
                "",
                "MATCHED SKILLS:",
                *[f"- {skill}" for skill in matched_skills],
                "",
                "POTENTIAL SKILL GAPS:",
                *[f"- {skill}" for skill in missing_skills],
                "",
                "Note: Skill matching is based on text detection. "
                "It does not prove proficiency or predict a "
                "confirmed hiring decision.",
            ]

            st.download_button(
                "Download Analysis Report",
                data="\n".join(report),
                file_name="resume_analysis.txt",
                mime="text/plain",
            )

        except Exception as error:
            st.error(f"Analysis failed: {error}")
            st.exception(error)
